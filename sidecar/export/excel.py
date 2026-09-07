"""sidecar/export/excel.py — Generator XLSX (openpyxl)"""
from __future__ import annotations
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def build_excel_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KPI Summary"
    ws["A1"] = "Cluster C1 \u2014 KPI Summary (TelecomAgent)"
    ws["A1"].font = Font(bold=True, color="7c3aed", size=14)
    ws["A2"] = "Generated: DT_Jakarta_C1_0409.csv  \u2022  142.3k rows  \u2022  847 cells"
    ws["A2"].font = Font(size=9, color="71717a", italic=True)
    data = [
        ["KPI", "Value", "Target", "Status"],
        ["RSRP \u2265 -100 dBm", "94.2%", "95%", "\u2717 below"],
        ["SINR \u2265 5 dB", "81.4%", "80%", "\u2713"],
        ["DL Throughput", "42.7 Mbps", "30 Mbps", "\u2713"],
        ["RSRP Avg", "-87.3 dBm", "\u2014", "\u2014"],
        ["SINR Avg", "7.2 dB", "\u2014", "\u2014"],
    ]
    thin = Side(style="thin", color="D4D4D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r, row in enumerate(data, start=4):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border
            cell.alignment = Alignment(horizontal="center" if c > 1 else "left", vertical="center")
            if r == 4:
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.fill = PatternFill("solid", fgColor="7c3aed")
            else:
                cell.font = Font(size=10)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws2 = wb.create_sheet("Worst Spots")
    ws2["A1"] = "Top 5 Worst Spots \u2014 RCA Engine"
    ws2["A1"].font = Font(bold=True, size=12, color="7c3aed")
    hdr = ["#", "Cell", "Issue", "RSRP", "SINR", "Rekomendasi"]
    for c, h in enumerate(hdr, 1):
        cell = ws2.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="7c3aed")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    rows = [
        [1, "JKT_1023_2", "Overshooting", "-108", "2.1", "downtilt 3\u00b0\u21925\u00b0"],
        [2, "JKT_1018_1", "PCI confusion", "-102", "3.4", "PCI 148\u2192312"],
        [3, "JKT_1015_1", "Missing neighbor", "-99", "4.0", "Add nbr \u21921022_2"],
        [4, "JKT_1040_3", "Weak coverage", "-110", "1.8", "tilt + azimuth"],
        [5, "JKT_1022_2", "HO fail", "-105", "2.5", "CIO tuning"],
    ]
    for r, row in enumerate(rows, start=4):
        for c, val in enumerate(row, 1):
            cell = ws2.cell(row=r, column=c, value=val)
            cell.font = Font(size=10)
            cell.alignment = Alignment(horizontal="center" if c != 2 else "left", vertical="center")
            cell.border = border
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["F"].width = 22
    ws3 = wb.create_sheet("Raw Sample")
    ws3["A1"] = "Sample rows (first 5) \u2014 DT_Jakarta_C1_0409.csv"
    ws3["A1"].font = Font(italic=True, size=9, color="71717a")
    sample_hdr = ["Time", "Lat", "Lon", "RSRP", "SINR", "DL Thr", "Cell"]
    for c, h in enumerate(sample_hdr, 1):
        cell = ws3.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor="27272a")
        cell.alignment = Alignment(horizontal="center")
    sample = [
        ["10:00:01", "-6.208", "106.845", "-92", "8.1", "45.2", "JKT_1023_2"],
        ["10:00:05", "-6.209", "106.846", "-88", "7.5", "52.1", "JKT_1023_2"],
        ["10:00:10", "-6.210", "106.847", "-108", "2.1", "5.3", "JKT_1023_2"],
        ["10:00:15", "-6.211", "106.848", "-85", "9.2", "61.0", "JKT_1018_1"],
        ["10:00:20", "-6.212", "106.849", "-99", "4.0", "22.4", "JKT_1015_1"],
    ]
    for r, row in enumerate(sample, start=4):
        for c, val in enumerate(row, 1):
            ws3.cell(row=r, column=c, value=val).font = Font(size=9)
    for c in range(1, 8):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
