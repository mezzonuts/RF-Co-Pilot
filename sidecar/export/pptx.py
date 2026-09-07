"""sidecar/export/pptx.py — Generator PPTX (python-pptx) 5 slides"""
from __future__ import annotations
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def build_pptx_bytes() -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    def bg(slide, rgb=(10, 10, 15)):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*rgb)

    def title_box(slide, text, sub=None):
        tx = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.7))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x7C, 0x3A, 0xED)
        if sub:
            tx2 = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12), Inches(0.4))
            tf2 = tx2.text_frame
            tf2.word_wrap = True
            p2 = tf2.paragraphs[0]
            p2.text = sub
            p2.font.size = Pt(11)
            p2.font.color.rgb = RGBColor(0xA1, 0xA1, 0xAA)

    def bullets(slide, items, left=0.6, top=1.6, w=12, h=5):
        tx = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        tf = tx.text_frame
        tf.word_wrap = True
        for i, b in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.level = 0
            p.font.size = Pt(13)
            p.font.color.rgb = RGBColor(0xE4, 0xE4, 0xE7)
            p.space_after = Pt(7)

    def kpi_table(slide):
        rows, cols = 4, 4
        tbl = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(1.7), Inches(12), Inches(2.2)).table
        hdr = ["KPI", "Value", "Target", "Status"]
        vals = [
            ["RSRP \u2265 -100 dBm", "94.2%", "95%", "\u2717 below"],
            ["SINR \u2265 5 dB", "81.4%", "80%", "\u2713"],
            ["DL Throughput", "42.7 Mbps", "30 Mbps", "\u2713"],
        ]
        for c, h in enumerate(hdr):
            cell = tbl.cell(0, c)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x7C, 0x3A, 0xED)
            for par in cell.text_frame.paragraphs:
                par.alignment = PP_ALIGN.CENTER
                run = par.runs[0]
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for r, row in enumerate(vals, 1):
            for c, v in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = v
                for par in cell.text_frame.paragraphs:
                    par.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
                    run = par.runs[0]
                    run.font.size = Pt(10)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_box(s, "Cluster C1 \u2014 Executive Summary", "DT_Jakarta_C1_0409.csv  \u2022  142.3k rows  \u2022  847 cells  \u2022  TelecomAgent RF Co-Pilot")
    bullets(s, [
        "RSRP \u2265 -100 dBm: 94.2%  (target 95% \u2014 slightly below)",
        "SINR \u2265 5 dB: 81.4%  (target 80% \u2713)",
        "Avg DL Throughput: 42.7 Mbps  \u2022  RSRP Avg -87.3 dBm  \u2022  SINR Avg 7.2 dB",
        "Next: 5 worst spots + RCA recommendations \u2192 slides 4\u20135",
    ], top=1.9)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_box(s, "KPI Overview")
    kpi_table(s)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_box(s, "Coverage Map \u2014 RSRP Distribution")
    bullets(s, [
        "Heatmap: RSRP distribution across Cluster Jabodetabek C1 (Folium/GeoJSON placeholder)",
        "Hotspot lemah: koridor Jl. Sudirman \u2013 Kuningan (RSRP -105 s.d. -110 dBm)",
        "Overshooting terdeteksi: JKT_1023_2 azimuth 120\u00b0 \u2192 perlu downtilt",
    ])
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_box(s, "Top 5 Worst Spots")
    bullets(s, [
        "1. JKT_1023_2 \u2014 RSRP -108 dBm / SINR 2.1 dB \u2014 Overshooting \u2192 downtilt 3\u00b0\u21925\u00b0",
        "2. JKT_1018_1 \u2194 1020_3 \u2014 PCI confusion 148\u2192312 \u2192 re-plan PCI",
        "3. JKT_1015_1 \u2192 1022_2 \u2014 Missing neighbor, 42 HO fails \u2192 Add neighbor",
        "4. JKT_1040_3 \u2014 Weak coverage -110 dBm \u2192 tilt + azimuth optimization",
        "5. JKT_1022_2 \u2014 HO fail cluster \u2192 CIO tuning",
    ])
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title_box(s, "Recommendations \u2014 RCA Engine")
    bullets(s, [
        "Downtilt JKT_1023_2 3\u00b0\u21925\u00b0 (priority 1) \u2014 estimasi +4% RSRP \u2265-100",
        "PCI re-plan 148 \u2192 312 untuk JKT_1018_1 / 1020_3 \u2014 cegah collision",
        "Add neighbor JKT_1015_1 \u2192 JKT_1022_2 \u2014 turunkan HO fail 42 \u2192 <5",
        "Follow-up Drive Test 7 hari setelah perubahan \u2014 validate KPI",
        "Export: Excel detail (3 sheets) + raw sample \u2192 D:\\TelecomReports\\",
    ])
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()
